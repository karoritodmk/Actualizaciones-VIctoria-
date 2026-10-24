import csv
import io
import re
import unicodedata
from datetime import date, datetime

from fastapi import HTTPException, UploadFile
from pydantic import ValidationError
from sqlalchemy.orm import Session

from models.ficha import Ficha
from models.rol import Rol
from models.usuario import Usuario
from schemas.usuario import UsuarioCreate
from services import usuario_service

EXTENSIONES_VALIDAS = (".csv", ".tsv", ".txt", ".xlsx")
TAMANO_MAXIMO_BYTES = 5 * 1024 * 1024
MAX_FILAS = 2000

# campo interno -> nombres de columna aceptados (ya normalizados)
ALIAS_COLUMNAS = {
    "nombre": {"nombre", "nombres", "name"},
    "apellido": {"apellido", "apellidos", "lastname"},
    "tipo_documento": {
        "tipo_documento", "tipodocumento", "tipo_de_documento", "tipo_doc", "tipodoc", "tipo",
    },
    "numero_documento": {
        "numero_documento", "numerodocumento", "numero_de_documento", "documento",
        "num_documento", "nro_documento", "n_documento", "identificacion", "cedula",
    },
    "fecha_nacimiento": {
        "fecha_nacimiento", "fechanacimiento", "fecha_de_nacimiento", "fecha_nac", "nacimiento",
    },
    "correo_personal": {
        "correo_personal", "correopersonal", "email_personal", "correo", "email", "e_mail",
    },
    "correo_sena": {
        "correo_sena", "correosena", "correo_institucional", "email_sena", "correo_misena",
    },
    "telefono": {"telefono", "celular", "tel", "movil", "telefono_celular"},
}
OBLIGATORIAS = ("nombre", "numero_documento", "fecha_nacimiento", "correo_personal", "correo_sena")

ETIQUETAS = {
    "nombre": "Nombre",
    "apellido": "Apellido",
    "tipo_documento": "Tipo de documento",
    "numero_documento": "Número de documento",
    "fecha_nacimiento": "Fecha de nacimiento",
    "correo_personal": "Correo personal",
    "correo_sena": "Correo SENA",
    "telefono": "Teléfono",
    "contrasena": "Contraseña",
}


def generar_password_temporal(numero_documento: str) -> str:
    """Contraseña inicial: 'Sena@' + documento. Cumple la política de seguridad
    (mayúscula, minúscula, número, especial, 8+). El aprendiz debe cambiarla."""
    clave = f"Sena@{numero_documento}"
    if not any(c.isdigit() for c in clave):
        clave += "1"
    return clave.ljust(8, "0")


def _normalizar(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def _texto(valor) -> str:
    """Convierte cualquier celda (texto, número, fecha de Excel) en texto limpio."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))  # Excel guarda 1016031850 como 1016031850.0
    return str(valor).strip()


def _decodificar(crudo: bytes) -> str:
    if crudo.startswith((b"\xff\xfe", b"\xfe\xff")):
        return crudo.decode("utf-16")
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return crudo.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return crudo.decode("latin-1")


def _filas_de_texto(crudo: bytes) -> list[list]:
    texto = _decodificar(crudo)
    primera = next((l for l in texto.splitlines() if l.strip()), "")
    separador = max([",", ";", "\t", "|"], key=primera.count)
    if primera.count(separador) == 0:
        separador = ","
    return list(csv.reader(io.StringIO(texto), delimiter=separador))


def _filas_de_xlsx(crudo: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(
            status_code=400,
            detail=(
                "Este servidor no puede leer archivos .xlsx todavía. "
                "Instala openpyxl (pip install openpyxl) o sube el listado como .csv."
            ),
        )
    try:
        libro = load_workbook(io.BytesIO(crudo), read_only=True, data_only=True)
        return [list(fila) for fila in libro.active.iter_rows(values_only=True)]
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo de Excel. ¿Está dañado?")


def _parsear_fecha(texto: str) -> date:
    texto = texto.strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError("Fecha de nacimiento no válida (usa AAAA-MM-DD o DD/MM/AAAA)")


def _mensaje_validacion(error: ValidationError) -> str:
    partes = []
    for e in error.errors():
        campo = str(e["loc"][-1]) if e.get("loc") else ""
        etiqueta = ETIQUETAS.get(campo, campo)
        tipo = e.get("type", "")
        if tipo == "missing":
            partes.append(f"{etiqueta} es obligatorio")
        elif campo in ("correo_personal", "correo_sena"):
            partes.append(f"{etiqueta} no es un correo válido")
        elif tipo == "value_error":
            partes.append(str(e.get("msg", "")).replace("Value error, ", ""))
        elif "too_long" in tipo:
            partes.append(f"{etiqueta} es demasiado largo")
        else:
            partes.append(f"{etiqueta} no es válido")
    return "; ".join(partes) or "Datos no válidos"


def procesar_archivo(db: Session, id_ficha: int, archivo: UploadFile) -> dict:
    nombre_archivo = (archivo.filename or "").lower()
    if not nombre_archivo.endswith(EXTENSIONES_VALIDAS):
        raise HTTPException(
            status_code=400,
            detail="Formato no válido. Sube un archivo .csv, .tsv, .txt o .xlsx.",
        )

    ficha = db.query(Ficha).filter(Ficha.id_ficha == id_ficha).first()
    if not ficha:
        raise HTTPException(status_code=404, detail=f"La ficha {id_ficha} no existe.")

    rol_aprendiz = db.query(Rol).filter(Rol.nombre_rol.ilike("aprendiz")).first()
    if not rol_aprendiz:
        raise HTTPException(status_code=500, detail="No existe el rol 'aprendiz'. Corre seed_roles.py primero.")

    crudo = archivo.file.read()
    if not crudo:
        raise HTTPException(status_code=400, detail="El archivo está vacío.")
    if len(crudo) > TAMANO_MAXIMO_BYTES:
        raise HTTPException(status_code=400, detail="El archivo pesa más de 5 MB.")

    filas = _filas_de_xlsx(crudo) if nombre_archivo.endswith(".xlsx") else _filas_de_texto(crudo)

    # --- encabezados ---
    indice_encabezado = next(
        (i for i, f in enumerate(filas) if any(_texto(c) for c in f)), None
    )
    if indice_encabezado is None:
        raise HTTPException(status_code=400, detail="El archivo no tiene datos.")

    columnas: dict[str, int] = {}
    for posicion, celda in enumerate(filas[indice_encabezado]):
        clave = _normalizar(celda)
        for campo, alias in ALIAS_COLUMNAS.items():
            if clave in alias and campo not in columnas:
                columnas[campo] = posicion
    faltan = [ETIQUETAS[c] for c in OBLIGATORIAS if c not in columnas]
    if faltan:
        raise HTTPException(
            status_code=400,
            detail=(
                "Faltan columnas obligatorias en la primera fila: " + ", ".join(faltan) + ". "
                "Columnas esperadas: nombre, apellido, tipo_documento, numero_documento, "
                "fecha_nacimiento, correo_personal, correo_sena, telefono."
            ),
        )

    datos_filas = filas[indice_encabezado + 1:]
    if len(datos_filas) > MAX_FILAS:
        raise HTTPException(status_code=400, detail=f"El archivo tiene más de {MAX_FILAS} filas.")

    creados: list[str] = []
    errores: list[dict] = []
    vistos_documento: set[str] = set()
    vistos_correo: set[str] = set()
    total = 0

    for desplazamiento, fila in enumerate(datos_filas, start=indice_encabezado + 2):
        if not any(_texto(c) for c in fila):
            continue  # fila en blanco
        total += 1
        try:
            def valor(campo: str) -> str:
                i = columnas.get(campo)
                return _texto(fila[i]) if i is not None and i < len(fila) else ""

            documento = valor("numero_documento")
            correo_personal = valor("correo_personal")
            correo_sena = valor("correo_sena")

            if not valor("nombre"):
                raise ValueError("El nombre es obligatorio")
            if not documento:
                raise ValueError("El número de documento es obligatorio")
            if documento in vistos_documento:
                raise ValueError(f"El documento {documento} está repetido en el archivo")
            if correo_personal.lower() in vistos_correo or correo_sena.lower() in vistos_correo:
                raise ValueError("Un correo está repetido en el archivo")

            nacimiento = _parsear_fecha(valor("fecha_nacimiento")) if valor("fecha_nacimiento") else None
            if nacimiento is None:
                raise ValueError("Fecha de nacimiento es obligatoria")

            nuevo = UsuarioCreate(
                nombre=valor("nombre"),
                apellido=valor("apellido") or None,
                id_rol=rol_aprendiz.id_rol,
                estado="activo",
                fecha_nacimiento=nacimiento,
                numero_documento=documento,
                tipo_documento=valor("tipo_documento") or "CC",
                telefono=valor("telefono") or None,
                correo_personal=correo_personal,
                correo_sena=correo_sena,
                id_ficha=id_ficha,
                contrasena=generar_password_temporal(documento),
            )

            ya_existe = (
                db.query(Usuario)
                .filter(
                    (Usuario.numero_documento == nuevo.numero_documento)
                    | (Usuario.correo_personal == str(nuevo.correo_personal))
                    | (Usuario.correo_sena == str(nuevo.correo_sena))
                )
                .first()
            )
            if ya_existe:
                if ya_existe.numero_documento == nuevo.numero_documento:
                    raise ValueError(f"Ya existe un usuario con el documento {nuevo.numero_documento}")
                raise ValueError("Ya existe un usuario con ese correo")

            usuario_service.crear_usuario(db, nuevo)
            vistos_documento.add(documento)
            vistos_correo.add(correo_personal.lower())
            vistos_correo.add(correo_sena.lower())
            creados.append(documento)
        except ValidationError as e:  # ValidationError hereda de ValueError: va primero
            db.rollback()
            errores.append({"fila": desplazamiento, "error": _mensaje_validacion(e)})
        except ValueError as e:
            db.rollback()
            errores.append({"fila": desplazamiento, "error": str(e)})
        except Exception:
            db.rollback()
            errores.append({"fila": desplazamiento, "error": "No se pudo guardar esta fila."})

    return {"creados": creados, "errores": errores, "total": total}